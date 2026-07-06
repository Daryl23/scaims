from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
SCREENSHOT_DIR = BASE_DIR / "screenshots_clean"

OUTPUTS = {
    "english": BASE_DIR / "SCAIMS_User_Manual_English.docx",
    "tagalog": BASE_DIR / "SCAIMS_User_Manual_Tagalog.docx",
}

SCREENSHOTS = {
    "home": "01_home_page_clean.png",
    "admin_login": "02_admin_login_clean.png",
    "farmer_login": "03_farmer_login_clean.png",
    "rsbsa": "04_rsbsa_verification_clean.png",
    "registration": "06_registration_form_clean.png",
    "account": "07_create_account_clean.png",
    "admin_users": "08_admin_user_management_clean.png",
    "admin_parcels": "09_admin_farm_parcel_clean.png",
    "admin_suggestions": "10_admin_suggestions_clean.png",
    "admin_crop_reports": "11_admin_crop_reports_clean.png",
    "admin_damage_reports": "12_admin_damage_reports_clean.png",
    "farmer_updates": "13_farmer_home_updates_clean.png",
    "report_form": "14_farmer_report_form_clean.png",
    "crop_fields": "15_farmer_crop_report_fields_clean.png",
    "damage_fields": "16_farmer_damage_report_fields_clean.png",
    "farm_parcel": "17_farmer_add_farm_parcel_clean.png",
    "post_updates": "18_admin_post_updates_clean.png",
    "admin_home": "19_admin_dashboard_home_clean.png",
    "reports_overview": "20_admin_reports_overview_clean.png",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="C9D3DF", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn("w:{}".format(m)))
        if node is None:
            node = OxmlElement("w:{}".format(m))
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.25):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def configure_document(doc, title):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Title", 20, "1F4D78", 0, 8),
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.space_after = Pt(3)
    bullet.paragraph_format.line_spacing = 1.15

    doc.core_properties.title = title
    doc.core_properties.author = "SCAIMS Documentation"

    footer = section.footer.paragraphs[0]
    footer.text = title
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(96, 108, 120)


def add_cover(doc, title, subtitle, updated_label, note):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(72)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(31, 77, 120)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(66, 78, 90)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(updated_label)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(96, 108, 120)

    doc.add_paragraph()
    add_callout(doc, note, fill="E8EEF5")
    doc.add_page_break()


def add_callout(doc, text, fill="F2F6FA"):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell)
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=0)
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(54, 67, 80)
    doc.add_paragraph()


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_screenshot(doc, screenshot_key, caption):
    path = SCREENSHOT_DIR / SCREENSHOTS[screenshot_key]
    if not path.exists():
        raise FileNotFoundError(path)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.15))

    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(0)
    caption_run = caption_p.add_run(caption)
    caption_run.italic = True
    caption_run.font.size = Pt(8.5)
    caption_run.font.color.rgb = RGBColor(96, 108, 120)


def add_step(doc, title, bullets, screenshot_key, caption, page_break=True):
    doc.add_heading(title, level=2)
    for bullet in bullets:
        add_bullet(doc, bullet)
    add_screenshot(doc, screenshot_key, caption)
    if page_break:
        doc.add_page_break()


def add_section_heading(doc, title, description):
    doc.add_heading(title, level=1)
    p = doc.add_paragraph(description)
    set_paragraph_spacing(p, after=8)


def add_intro(doc, heading, purpose, before_title, before_items, reminder):
    doc.add_heading(heading, level=1)
    p = doc.add_paragraph(purpose)
    set_paragraph_spacing(p)

    doc.add_heading(before_title, level=2)
    for item in before_items:
        add_bullet(doc, item)
    add_callout(doc, reminder, fill="FFF4D8")


ENGLISH_STEPS = [
    {
        "title": "Step 1. Open the SCAIMS home page",
        "bullets": [
            "Use this page as the starting point of the system.",
            "Choose Register Your Farm if you are a new farmer and still need an account.",
            "Choose Login if you already have an account.",
            "Choose View Crop Map only when you want to look at the crop map.",
        ],
        "screenshot": "home",
        "caption": "SCAIMS home page.",
    },
    {
        "title": "Step 2. Log in as a farmer",
        "bullets": [
            "Enter your Username and Password.",
            "Complete the reCAPTCHA check if it appears.",
            "Click Sign In as Farmer.",
            "If you do not have an account yet, click Create a Farmer Account.",
        ],
        "screenshot": "farmer_login",
        "caption": "Farmer login page.",
    },
    {
        "title": "Step 3. Answer the RSBSA question",
        "bullets": [
            "Read the question asking if you are already registered in RSBSA.",
            "If you are already registered, click Oo, ako ay nakarehistro.",
            "If you are not yet registered, click Hindi, kailangan ko pang magparehistro and follow the RSBSA instructions first.",
            "Return to this registration page after your RSBSA registration is approved.",
        ],
        "screenshot": "rsbsa",
        "caption": "RSBSA verification page.",
    },
    {
        "title": "Step 4. Fill out the farmer registration form",
        "bullets": [
            "Enter your name, birthdate, sex, contact number, address, barangay, and RSBSA Number.",
            "Use this RSBSA Number format: XX-XX-XX-XXX-XXXXXX.",
            "Use an 11-digit Philippine mobile number that starts with 09.",
            "Read the terms and privacy notice, check the box if you agree, then click Submit.",
        ],
        "screenshot": "registration",
        "caption": "Farmer registration form.",
    },
    {
        "title": "Step 5. Create your farmer account",
        "bullets": [
            "Choose a username that you can remember.",
            "Create a strong password with at least 8 characters, including uppercase, lowercase, number, and special character.",
            "Type the same password again in Confirm Password.",
            "Upload a clear profile image, then click Create Account.",
            "Write down your username and password in a safe place.",
        ],
        "screenshot": "account",
        "caption": "Create account page.",
    },
    {
        "title": "Step 6. Read the latest updates",
        "bullets": [
            "After logging in, open Home from the left menu.",
            "Read the latest announcements from the agriculture office.",
            "Check this page before submitting reports so you do not miss important reminders.",
        ],
        "screenshot": "farmer_updates",
        "caption": "Farmer home page with latest updates.",
    },
    {
        "title": "Step 7. Open the report form",
        "bullets": [
            "Click Report Forms from the left menu.",
            "Under Choose Type of Report, select Crop Harvest Report for harvest information.",
            "Select Crop Damage Report if crops were damaged by pest, disease, weather, or another cause.",
        ],
        "screenshot": "report_form",
        "caption": "Report form page before choosing a report type.",
    },
    {
        "title": "Step 8. Submit a crop harvest report",
        "bullets": [
            "Choose the parcel location and crop details.",
            "Enter the planting date, harvest date, estimated harvest in kilograms, crop status, number of trees if applicable, and farm gate price.",
            "For Number of Trees, type N/A if the crop does not use tree count.",
            "Coordinates are optional. If you are not sure, choose Do not include coordinates or leave them blank.",
            "Review the details, then click Submit Report.",
        ],
        "screenshot": "crop_fields",
        "caption": "Crop harvest report fields.",
    },
    {
        "title": "Step 9. Submit a crop damage report",
        "bullets": [
            "Choose the parcel location and crop details.",
            "Enter the date of occurrence, cause of damage, percentage of damage, estimated loss in kilograms, number of trees if applicable, and farm gate price.",
            "For Number of Trees, type N/A if the crop does not use tree count.",
            "Coordinates are optional. If you are not sure, choose Do not include coordinates or leave them blank.",
            "Review the details, then click Submit Report.",
        ],
        "screenshot": "damage_fields",
        "caption": "Crop damage report fields.",
    },
    {
        "title": "Step 10. Add a farm parcel",
        "bullets": [
            "Click Farm Parcel from the left menu.",
            "Enter the parcel location, parcel size in hectares, ownership type, number of crops, farm type, and whether you are an organic practitioner.",
            "If more crop fields appear after entering the number of crops, fill them out before saving.",
            "Click Add Parcel when the information is complete.",
        ],
        "screenshot": "farm_parcel",
        "caption": "Add farm parcel form.",
    },
    {
        "title": "Step 11. Log out of the farmer account",
        "bullets": [
            "When you are finished, click Logout from the left menu.",
            "Logging out helps protect your account, especially when using a shared computer or phone.",
        ],
        "screenshot": "farmer_updates",
        "caption": "Farmer menu with Logout option.",
    },
]


ENGLISH_ADMIN_STEPS = [
    {
        "title": "Step 12. Log in as an admin",
        "bullets": [
            "Enter the Admin Username and Password provided by the office.",
            "Complete the reCAPTCHA check if it appears.",
            "Click Sign in as Admin.",
        ],
        "screenshot": "admin_login",
        "caption": "Admin login page.",
    },
    {
        "title": "Step 13. Check the admin dashboard",
        "bullets": [
            "Use the dashboard to see totals for registrations, farm parcels, and submitted reports.",
            "Click a card arrow or use the left menu to open the related section.",
            "Review this page first to get a quick summary of system activity.",
        ],
        "screenshot": "admin_home",
        "caption": "Admin dashboard summary.",
    },
    {
        "title": "Step 14. Manage user records",
        "bullets": [
            "Open User Management from the left menu.",
            "Use the search box or filters to find a farmer or admin record.",
            "Use Edit when information must be corrected. The system may ask for the admin PIN before saving changes.",
            "Use Deactivate or Activate when an account status needs to be changed.",
            "Use Export if the office needs a copy of the list.",
        ],
        "screenshot": "admin_users",
        "caption": "Admin user management page.",
    },
    {
        "title": "Step 15. Review farm parcels",
        "bullets": [
            "Open Farm Parcel from the left menu.",
            "Search by farmer name or use the filters for location, crop commodity, or crop type.",
            "Click View Parcel Details to check the parcels registered under a farmer.",
            "Use Export if the office needs a copy of the records.",
        ],
        "screenshot": "admin_parcels",
        "caption": "Admin farm parcel records.",
    },
    {
        "title": "Step 16. Post updates for farmers",
        "bullets": [
            "Open Post Updates from the left menu.",
            "Type the announcement or reminder in the text box.",
            "Click Add Images if a picture should be included.",
            "Click Post Update. The post will appear in the farmers' Latest Updates page.",
            "Delete only posts that are outdated or incorrect.",
        ],
        "screenshot": "post_updates",
        "caption": "Admin post updates page.",
    },
    {
        "title": "Step 17. Review farmer suggestions",
        "bullets": [
            "Open Suggestions from the left menu.",
            "Check New Suggestions first, then review Previous Suggestions if needed.",
            "Use the month filter to narrow the list.",
            "Mark suggestions as read after review, and delete only when the office no longer needs the record.",
            "Use Export Suggestions if a copy is needed.",
        ],
        "screenshot": "admin_suggestions",
        "caption": "Admin suggestions page.",
    },
    {
        "title": "Step 18. Open reports overview or submit for a farmer",
        "bullets": [
            "Open Reports from the left menu.",
            "Use the Crop Reports card for harvest reports and the Damage Reports card for damage reports.",
            "If a farmer needs help, use Submit Report For A Farmer, search for the farmer, choose the report type, and complete the form on their behalf.",
            "Review charts to see barangays or crops that need attention.",
        ],
        "screenshot": "reports_overview",
        "caption": "Admin reports overview page.",
    },
    {
        "title": "Step 19. Review crop harvest reports",
        "bullets": [
            "Open Crop Reports from the reports section.",
            "Use search and filters for farmer, reporter, month, crop, or barangay.",
            "Click View to open the report details.",
            "If coordinates are included, use the GIS map options to view the location or directions.",
            "Use Export or View Graph when the office needs a report summary.",
        ],
        "screenshot": "admin_crop_reports",
        "caption": "Admin crop harvest reports page.",
    },
    {
        "title": "Step 20. Review crop damage reports",
        "bullets": [
            "Open Damage Reports from the reports section.",
            "Use search and filters for farmer, reporter, date, crop, or barangay.",
            "Click View to open the report details.",
            "If coordinates are included, use the GIS map options to view the location or directions.",
            "Use Export or View Graph when the office needs a report summary.",
        ],
        "screenshot": "admin_damage_reports",
        "caption": "Admin crop damage reports page.",
    },
    {
        "title": "Step 21. Log out of the admin account",
        "bullets": [
            "When office work is finished, click Logout from the left menu.",
            "Logging out protects farmer records and office data, especially on shared computers.",
        ],
        "screenshot": "admin_home",
        "caption": "Admin menu with Logout option.",
    },
]


TAGALOG_STEPS = [
    {
        "title": "Hakbang 1. Buksan ang SCAIMS home page",
        "bullets": [
            "Ito ang unang pahina ng system.",
            "Kung bago ka pa lang gagamit at wala pang account, pindutin ang Register Your Farm.",
            "Kung mayroon ka nang account, pindutin ang Login.",
            "Pindutin lang ang View Crop Map kung gusto mong tingnan ang mapa ng mga pananim.",
        ],
        "screenshot": "home",
        "caption": "SCAIMS home page.",
    },
    {
        "title": "Hakbang 2. Mag-login bilang magsasaka",
        "bullets": [
            "Ilagay ang iyong Username at Password.",
            "Sagutan ang reCAPTCHA kung may lumabas na check.",
            "Pindutin ang Sign In as Farmer.",
            "Kung wala ka pang account, pindutin ang Create a Farmer Account.",
        ],
        "screenshot": "farmer_login",
        "caption": "Login page ng magsasaka.",
    },
    {
        "title": "Hakbang 3. Sagutin ang tanong tungkol sa RSBSA",
        "bullets": [
            "Basahin ang tanong kung ikaw ay rehistrado na sa RSBSA.",
            "Kung rehistrado ka na, pindutin ang Oo, ako ay nakarehistro.",
            "Kung hindi ka pa rehistrado, pindutin ang Hindi, kailangan ko pang magparehistro at sundin muna ang gabay para sa RSBSA.",
            "Bumalik sa registration page kapag aprubado na ang iyong RSBSA registration.",
        ],
        "screenshot": "rsbsa",
        "caption": "RSBSA verification page.",
    },
    {
        "title": "Hakbang 4. Sagutan ang registration form ng magsasaka",
        "bullets": [
            "Ilagay ang pangalan, kaarawan, kasarian, contact number, tirahan, barangay, at RSBSA Number.",
            "Sundin ang format ng RSBSA Number: XX-XX-XX-XXX-XXXXXX.",
            "Para sa contact number, gumamit ng 11-digit mobile number na nagsisimula sa 09.",
            "Basahin ang terms at privacy notice. Kung sumasang-ayon ka, lagyan ng check ang kahon at pindutin ang Submit.",
        ],
        "screenshot": "registration",
        "caption": "Registration form ng magsasaka.",
    },
    {
        "title": "Hakbang 5. Gumawa ng account ng magsasaka",
        "bullets": [
            "Pumili ng username na madali mong matatandaan.",
            "Gumawa ng matibay na password na may hindi bababa sa 8 character, may malaking titik, maliit na titik, numero, at special character.",
            "I-type muli ang parehong password sa Confirm Password.",
            "Mag-upload ng malinaw na larawan para sa profile, pagkatapos ay pindutin ang Create Account.",
            "Itago sa ligtas na lugar ang iyong username at password.",
        ],
        "screenshot": "account",
        "caption": "Create account page.",
    },
    {
        "title": "Hakbang 6. Basahin ang latest updates",
        "bullets": [
            "Pagkatapos mag-login, buksan ang Home sa kaliwang menu.",
            "Basahin ang mga anunsyo mula sa tanggapan ng agrikultura.",
            "Tingnan muna ang pahinang ito bago magpasa ng ulat para wala kang mamiss na mahalagang paalala.",
        ],
        "screenshot": "farmer_updates",
        "caption": "Home page ng magsasaka na may latest updates.",
    },
    {
        "title": "Hakbang 7. Buksan ang report form",
        "bullets": [
            "Pindutin ang Report Forms sa kaliwang menu.",
            "Sa Choose Type of Report, piliin ang Crop Harvest Report kung ulat ito tungkol sa ani.",
            "Piliin ang Crop Damage Report kung may nasirang pananim dahil sa peste, sakit, panahon, o ibang dahilan.",
        ],
        "screenshot": "report_form",
        "caption": "Report form page bago pumili ng uri ng ulat.",
    },
    {
        "title": "Hakbang 8. Magpasa ng ulat tungkol sa ani",
        "bullets": [
            "Piliin ang lokasyon ng parcel at detalye ng pananim.",
            "Ilagay ang petsa ng tanim, petsa ng ani, tantiyang ani sa kilo, kalagayan ng pananim, bilang ng puno kung angkop, at farm gate price.",
            "Sa Number of Trees, ilagay ang N/A kung hindi ito ginagamit sa pananim mo.",
            "Optional ang coordinates. Kung hindi ka sigurado, piliin ang Do not include coordinates o iwan itong blanko.",
            "Suriin muna ang detalye, pagkatapos ay pindutin ang Submit Report.",
        ],
        "screenshot": "crop_fields",
        "caption": "Mga field para sa ulat tungkol sa ani.",
    },
    {
        "title": "Hakbang 9. Magpasa ng ulat tungkol sa nasirang pananim",
        "bullets": [
            "Piliin ang lokasyon ng parcel at detalye ng pananim.",
            "Ilagay ang petsa ng pangyayari, sanhi ng pinsala, porsyento ng pinsala, tantiyang nawala sa kilo, bilang ng puno kung angkop, at farm gate price.",
            "Sa Number of Trees, ilagay ang N/A kung hindi ito ginagamit sa pananim mo.",
            "Optional ang coordinates. Kung hindi ka sigurado, piliin ang Do not include coordinates o iwan itong blanko.",
            "Suriin muna ang detalye, pagkatapos ay pindutin ang Submit Report.",
        ],
        "screenshot": "damage_fields",
        "caption": "Mga field para sa ulat tungkol sa nasirang pananim.",
    },
    {
        "title": "Hakbang 10. Magdagdag ng farm parcel",
        "bullets": [
            "Pindutin ang Farm Parcel sa kaliwang menu.",
            "Ilagay ang lokasyon ng parcel, sukat sa ektarya, uri ng pagmamay-ari, bilang ng pananim, farm type, at kung ikaw ay organic practitioner.",
            "Kung may dagdag na crop fields na lumabas pagkatapos ilagay ang bilang ng pananim, sagutan muna ang mga ito bago i-save.",
            "Pindutin ang Add Parcel kapag kumpleto na ang impormasyon.",
        ],
        "screenshot": "farm_parcel",
        "caption": "Add farm parcel form.",
    },
    {
        "title": "Hakbang 11. Mag-logout sa account ng magsasaka",
        "bullets": [
            "Kapag tapos ka na, pindutin ang Logout sa kaliwang menu.",
            "Mahalagang mag-logout lalo na kung gamit mo ay shared na computer o cellphone.",
        ],
        "screenshot": "farmer_updates",
        "caption": "Menu ng magsasaka na may Logout.",
    },
]


TAGALOG_ADMIN_STEPS = [
    {
        "title": "Hakbang 12. Mag-login bilang admin",
        "bullets": [
            "Ilagay ang Admin Username at Password na ibinigay ng opisina.",
            "Sagutan ang reCAPTCHA kung may lumabas na check.",
            "Pindutin ang Sign in as Admin.",
        ],
        "screenshot": "admin_login",
        "caption": "Admin login page.",
    },
    {
        "title": "Hakbang 13. Tingnan ang admin dashboard",
        "bullets": [
            "Sa dashboard makikita ang kabuuang bilang ng registrations, farm parcels, at submitted reports.",
            "Pindutin ang arrow sa card o gamitin ang kaliwang menu para buksan ang kaugnay na bahagi.",
            "Unahin itong tingnan para mabilis makita ang kalagayan ng system.",
        ],
        "screenshot": "admin_home",
        "caption": "Admin dashboard summary.",
    },
    {
        "title": "Hakbang 14. Ayusin ang user records",
        "bullets": [
            "Buksan ang User Management sa kaliwang menu.",
            "Gamitin ang search box o filters para hanapin ang record ng magsasaka o admin.",
            "Gamitin ang Edit kung may kailangang itama. Maaaring humingi ang system ng admin PIN bago ma-save ang pagbabago.",
            "Gamitin ang Deactivate o Activate kung kailangang baguhin ang status ng account.",
            "Gamitin ang Export kung kailangan ng opisina ng kopya ng listahan.",
        ],
        "screenshot": "admin_users",
        "caption": "Admin user management page.",
    },
    {
        "title": "Hakbang 15. Suriin ang mga farm parcel",
        "bullets": [
            "Buksan ang Farm Parcel sa kaliwang menu.",
            "Mag-search gamit ang pangalan ng magsasaka o gumamit ng filters para sa location, crop commodity, o crop type.",
            "Pindutin ang View Parcel Details para makita ang parcels na nakarehistro sa magsasaka.",
            "Gamitin ang Export kung kailangan ng opisina ng kopya ng records.",
        ],
        "screenshot": "admin_parcels",
        "caption": "Admin farm parcel records.",
    },
    {
        "title": "Hakbang 16. Mag-post ng updates para sa mga magsasaka",
        "bullets": [
            "Buksan ang Post Updates sa kaliwang menu.",
            "I-type ang anunsyo o paalala sa text box.",
            "Pindutin ang Add Images kung may kasamang larawan.",
            "Pindutin ang Post Update. Makikita ito ng mga magsasaka sa kanilang Latest Updates page.",
            "Burahin lang ang post kung luma na o mali ang impormasyon.",
        ],
        "screenshot": "post_updates",
        "caption": "Admin post updates page.",
    },
    {
        "title": "Hakbang 17. Tingnan ang suggestions ng mga magsasaka",
        "bullets": [
            "Buksan ang Suggestions sa kaliwang menu.",
            "Unahing tingnan ang New Suggestions, pagkatapos ay tingnan ang Previous Suggestions kung kailangan.",
            "Gamitin ang month filter para mas madaling hanapin ang record.",
            "I-mark as read kapag nasuri na, at mag-delete lang kung hindi na kailangan ng opisina ang record.",
            "Gamitin ang Export Suggestions kung kailangan ng kopya.",
        ],
        "screenshot": "admin_suggestions",
        "caption": "Admin suggestions page.",
    },
    {
        "title": "Hakbang 18. Buksan ang reports overview o mag-submit ng ulat para sa magsasaka",
        "bullets": [
            "Buksan ang Reports sa kaliwang menu.",
            "Gamitin ang Crop Reports card para sa ulat ng ani at Damage Reports card para sa ulat ng pinsala.",
            "Kung kailangan ng tulong ng magsasaka, gamitin ang Submit Report For A Farmer, hanapin ang magsasaka, piliin ang report type, at kumpletuhin ang form para sa kanya.",
            "Tingnan ang charts para makita ang barangay o pananim na nangangailangan ng pansin.",
        ],
        "screenshot": "reports_overview",
        "caption": "Admin reports overview page.",
    },
    {
        "title": "Hakbang 19. Suriin ang crop harvest reports",
        "bullets": [
            "Buksan ang Crop Reports sa reports section.",
            "Gamitin ang search at filters para sa magsasaka, reporter, buwan, crop, o barangay.",
            "Pindutin ang View para buksan ang detalye ng report.",
            "Kung may coordinates, gamitin ang GIS map options para makita ang lokasyon o direksyon.",
            "Gamitin ang Export o View Graph kung kailangan ng opisina ng report summary.",
        ],
        "screenshot": "admin_crop_reports",
        "caption": "Admin crop harvest reports page.",
    },
    {
        "title": "Hakbang 20. Suriin ang crop damage reports",
        "bullets": [
            "Buksan ang Damage Reports sa reports section.",
            "Gamitin ang search at filters para sa magsasaka, reporter, petsa, crop, o barangay.",
            "Pindutin ang View para buksan ang detalye ng report.",
            "Kung may coordinates, gamitin ang GIS map options para makita ang lokasyon o direksyon.",
            "Gamitin ang Export o View Graph kung kailangan ng opisina ng report summary.",
        ],
        "screenshot": "admin_damage_reports",
        "caption": "Admin crop damage reports page.",
    },
    {
        "title": "Hakbang 21. Mag-logout sa admin account",
        "bullets": [
            "Kapag tapos na ang gawain sa opisina, pindutin ang Logout sa kaliwang menu.",
            "Mahalagang mag-logout para maprotektahan ang records ng mga magsasaka at datos ng opisina.",
        ],
        "screenshot": "admin_home",
        "caption": "Admin menu na may Logout.",
    },
]


def build_manual(language):
    if language == "english":
        title = "SCAIMS / APGORRS User Manual"
        subtitle = "Step-by-step guide for farmers and agriculture office staff"
        updated_label = "Updated: July 1, 2026"
        cover_note = (
            "This manual is written for non-technical users. Some screens show the APGORRS "
            "name; those screens are part of the same SCAIMS agriculture system."
        )
        intro = {
            "heading": "How to Use This Manual",
            "purpose": (
                "Follow the steps in order. Compare your screen with the screenshot under each step. "
                "If the screen looks different, check that you opened the correct menu item before continuing."
            ),
            "before_title": "Before You Start",
            "before_items": [
                "Use a phone, tablet, laptop, or desktop with internet access.",
                "Prepare your username and password if you already have an account.",
                "Prepare your RSBSA Number if you are registering as a farmer.",
                "Prepare a clear profile image for account creation.",
                "Ask the agriculture office for help if you are unsure about RSBSA, parcel size, crop details, or report details.",
            ],
            "reminder": (
                "Important: Keep your username and password private. Do not share your account with another person."
            ),
            "farmer_title": "Part 1. Farmer Guide",
            "farmer_desc": "Use this section to register, log in, add farm parcels, read updates, and submit crop reports.",
            "admin_title": "Part 2. Admin and Office Staff Guide",
            "admin_desc": "Use this section to manage users, farm parcels, updates, suggestions, and submitted reports.",
        }
        farmer_steps = ENGLISH_STEPS
        admin_steps = ENGLISH_ADMIN_STEPS
        output = OUTPUTS["english"]
    else:
        title = "Manwal sa Paggamit ng SCAIMS / APGORRS"
        subtitle = "Madaling gabay para sa mga magsasaka at kawani ng Tanggapan ng Agrikultura"
        updated_label = "Na-update: July 1, 2026"
        cover_note = (
            "Ginawa ang manwal na ito para madaling sundan kahit hindi sanay sa computer. "
            "May ilang screen na may pangalang APGORRS; bahagi pa rin ito ng SCAIMS agriculture system."
        )
        intro = {
            "heading": "Paano Gamitin ang Manwal na Ito",
            "purpose": (
                "Sundin ang mga hakbang ayon sa pagkakasunod-sunod. Ihambing ang nakikita mo sa screen "
                "sa screenshot sa bawat hakbang. Kung iba ang itsura, siguraduhing tama ang menu na binuksan mo."
            ),
            "before_title": "Ihanda Bago Magsimula",
            "before_items": [
                "Gumamit ng cellphone, tablet, laptop, o desktop na may internet.",
                "Ihanda ang username at password kung may account ka na.",
                "Ihanda ang RSBSA Number kung magrerehistro bilang magsasaka.",
                "Maghanda ng malinaw na larawan para sa profile sa paggawa ng account.",
                "Magtanong sa Tanggapan ng Agrikultura kung hindi sigurado sa RSBSA, sukat ng sakahan, detalye ng pananim, o detalye ng ulat.",
            ],
            "reminder": (
                "Mahalaga: Itago nang maayos ang username at password. Huwag itong ipagamit sa ibang tao."
            ),
            "farmer_title": "Bahagi 1. Gabay para sa Magsasaka",
            "farmer_desc": "Gamitin ang bahaging ito para magparehistro, mag-login, magdagdag ng farm parcel, magbasa ng updates, at magpasa ng mga ulat tungkol sa pananim.",
            "admin_title": "Bahagi 2. Gabay para sa Admin at Kawani ng Opisina",
            "admin_desc": "Gamitin ang bahaging ito para mag-manage ng users, farm parcels, updates, suggestions, at submitted reports.",
        }
        farmer_steps = TAGALOG_STEPS
        admin_steps = TAGALOG_ADMIN_STEPS
        output = OUTPUTS["tagalog"]

    doc = Document()
    configure_document(doc, title)
    add_cover(doc, title, subtitle, updated_label, cover_note)
    add_intro(
        doc,
        intro["heading"],
        intro["purpose"],
        intro["before_title"],
        intro["before_items"],
        intro["reminder"],
    )
    doc.add_page_break()

    add_section_heading(doc, intro["farmer_title"], intro["farmer_desc"])
    for step in farmer_steps:
        add_step(doc, step["title"], step["bullets"], step["screenshot"], step["caption"])

    add_section_heading(doc, intro["admin_title"], intro["admin_desc"])
    for idx, step in enumerate(admin_steps):
        add_step(
            doc,
            step["title"],
            step["bullets"],
            step["screenshot"],
            step["caption"],
            page_break=idx < len(admin_steps) - 1,
        )

    doc.save(output)
    return output


def main():
    for language in ("english", "tagalog"):
        output = build_manual(language)
        print(output)


if __name__ == "__main__":
    main()
